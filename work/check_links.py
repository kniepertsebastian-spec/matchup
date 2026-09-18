from pathlib import Path
import openpyxl, json, zipfile
from html.parser import HTMLParser
base=Path('outputs/Olaf_Top_26.18_Revision').resolve()
local_links=0;source_links=0;books=0
for f in base.rglob('*.xlsx'):
    w=openpyxl.load_workbook(f,data_only=False,read_only=False)
    books+=1
    for s in w:
        for row in s:
            for c in row:
                if c.hyperlink:
                    target=c.hyperlink.target
                    if target.startswith('https://'):source_links+=1
                    else:
                        assert (f.parent/target).is_file(),(f.name,target)
                        local_links+=1
                assert c.data_type!='e',(f.name,c.coordinate,c.value)
    if f.name=='Olaf_Top_Sammlung.xlsx':
        assert w['Übersicht'].freeze_panes=='B6'
        assert w['Alle Details'].freeze_panes=='C6'
        assert len(w['Übersicht'].tables)==1
        assert w['Übersicht'].max_row>=56
    w.close()
assert books==52 and local_links==51
class Links(HTMLParser):
    def __init__(self):super().__init__();self.targets=[];self.cards=0
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag=='a':self.targets.append(a['href'])
        if tag=='article' and a.get('class')=='card':self.cards+=1
p=Links();p.feed((base/'START.html').read_text(encoding='utf-8'))
assert p.cards==51
for target in p.targets:assert (base/target).is_file(),target
with zipfile.ZipFile(base.parent/'Olaf_Top_26.18_Revision_Komplett.zip') as z:
    assert len(z.namelist())==54
    assert not any('inspect' in n for n in z.namelist())
print(json.dumps({'readable_workbooks':books,'excel_local_links':local_links,'excel_source_links':source_links,'catalog_cards':p.cards,'catalog_links':len(p.targets),'zip_files':54},indent=2))
