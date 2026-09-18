from pathlib import Path
import json,zipfile,re,shutil,xml.etree.ElementTree as ET
import openpyxl
from PIL import Image,ImageDraw
root=Path.cwd().resolve();work=root/'work';out=root/'outputs/Warwick_Top_26.18';web=root/'olaf-local';downloads=web/'dist/downloads'
d=json.loads((web/'dist/warwick.json').read_text(encoding='utf-8'))
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
# Reuse the prior missing-feature adapter, limited to native hyperlinks.
source=(work/'package.py').read_text(encoding='utf-8')
helper=source[source.index('def add_native_links(f):'):source.index("for f in out.rglob('*.xlsx'):add_native_links(f)")]
helper=helper.replace("v.startswith('https://')", "(v.startswith('https://') and ' | ' not in v)").replace('Matchups/Olaf_vs_','Matchups/Warwick_vs_')
exec(helper)
for f in out.rglob('*.xlsx'):
 add_native_links(f)
 shutil.copy2(f,downloads/f.relative_to(out))
for base in [out,downloads]:
 for f in base.rglob('*.inspect.ndjson'):
  assert f.resolve().is_relative_to(root)
  target=work/'inspection-warwick'/('output' if base==out else 'web')/f.relative_to(base)
  target.parent.mkdir(parents=True,exist_ok=True);shutil.move(str(f),str(target))
for m in d['matchups']:
 f=out/'Matchups'/m['file'];w=openpyxl.load_workbook(f,data_only=True);s=w.active
 assert s['B6'].value==m['name'] and s['B7'].value==m['tier'] and s['B8'].value==m['rune'] and s['B9'].value==m['sums']
 labels=[s.cell(i,1).value for i in range(1,s.max_row+1)]
 nums={int(re.match(r'(\d+) ·',str(x)).group(1)) for x in labels if re.match(r'(\d+) ·',str(x))}
 assert nums==set(range(1,21)),f
 for row in s:
  for c in row:assert c.data_type!='e',(f,c.coordinate)
 w.close()
w=openpyxl.load_workbook(out/'Warwick_Top_Sammlung.xlsx',data_only=True)
assert len(w.worksheets)==5
for i,m in enumerate(d['matchups'],6):
 assert w['Übersicht'].cell(i,2).value==m['tier']
 assert w['Übersicht'].cell(i,9).hyperlink.target=='Matchups/'+m['file']
 assert (out/w['Übersicht'].cell(i,9).hyperlink.target).is_file()
 if m['stats']['n'] is None:assert w['Statistik'].cell(i,2).value is None
w.close()
pre=work/'previews-warwick';contacts=work/'contacts-warwick';contacts.mkdir(exist_ok=True)
for start in range(0,51,6):
 canvas=Image.new('RGB',(1200,1200),'#d8e1eb');draw=ImageDraw.Draw(canvas)
 for j,m in enumerate(d['matchups'][start:start+6]):
  im=Image.open(pre/(m['slug']+'.png')).convert('RGB');im.thumbnail((580,365))
  x=(j%2)*600;y=(j//2)*400;draw.text((x+10,y+5),m['name'],fill='#243b61');canvas.paste(im,(x+10,y+28))
 canvas.save(contacts/f'{start//6+1}.png')
archive=root/'outputs/Olaf_Warwick_Matchup_Webapp_Docker.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for f in web.rglob('*'):
  if f.is_file() and not any(x in f.parts for x in ['node_modules','.sites-runtime']):z.write(f,Path('olaf-local')/f.relative_to(web))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None;assert not any('inspect.ndjson' in n for n in z.namelist())
print(json.dumps({'new_workbooks':52,'complete_fields':20,'working_workbook_links':51,'contact_sheets':9,'zip_bytes':archive.stat().st_size}))
